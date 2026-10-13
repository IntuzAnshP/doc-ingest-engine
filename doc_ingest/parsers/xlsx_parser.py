import openpyxl
from doc_ingest.parsers.base import BaseParser
from doc_ingest.schemas.source_item import SourceItem
from doc_ingest.schemas.extracted_document import ExtractedDocument

class XLSXParser(BaseParser):
    def supports(self, content_type: str) -> bool:
        return content_type == "xlsx"

    def parse(self, source_item: SourceItem) -> ExtractedDocument:
        wb = openpyxl.load_workbook(source_item.raw_path, data_only=True)
        
        full_text = []
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            for row in sheet.iter_rows(values_only=True):
                row_values = [str(cell).strip() for cell in row if cell is not None and str(cell).strip()]
                if row_values:
                    row_str = " | ".join(row_values)
                    full_text.append(row_str)
            
        parser_metadata = {
            "sheet_names": wb.sheetnames
        }
        
        return ExtractedDocument(
            source_item=source_item,
            raw_text="\n".join(full_text),
            parser_metadata=parser_metadata
        )
